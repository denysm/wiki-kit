"""Extract a document into a Markdown raw source with frontmatter.

Usage: doc2md.py <plan.json>
plan.json: [{"src": "/abs/file", "out": "/abs/raw/.../name.md", "title": "...",
             "type": "notes|articles", "tags": [...], "source_date": "YYYY-MM-DD",
             "authors": [...], "source": "optional override", "extra": {"k": "v"}}]
Supports .pptx (slide text + notes), .docx (paragraphs, headings, tables),
.pdf (page text), .md/.txt (verbatim). Refuses to overwrite.
"""
import json
import re
import sys
from pathlib import Path

TODAY = __import__("datetime").date.today().isoformat()
HOME = str(Path.home())


def q(s: str) -> str:
    return '"' + str(s).replace("\\", "\\\\").replace('"', '\\"') + '"'


def pptx_md(p: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(p))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = "\n".join(par.text for par in shape.text_frame.paragraphs if par.text.strip())
                if t.strip():
                    texts.append(t.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [" | ".join(c.text.strip().replace("\n", " ") for c in r.cells) for r in shape.table.rows]
                texts.append("\n".join(f"| {r} |" for r in rows))
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        out.append(f"## Slide {i}\n\n" + ("\n\n".join(texts) or "*(no text)*"))
        if notes:
            out.append(f"**Speaker notes:**\n\n{notes}")
    return "\n\n".join(out)


def docx_md(p: Path) -> str:
    import docx
    d = docx.Document(str(p))
    out = []
    for block in d.element.body.iterchildren():
        tag = block.tag.split("}")[-1]
        if tag == "p":
            par = docx.text.paragraph.Paragraph(block, d)
            t = par.text.strip()
            if not t:
                continue
            style = (par.style.name or "").lower() if par.style is not None else ""
            m = re.match(r"heading (\d)", style)
            if m:
                out.append("#" * min(int(m.group(1)) + 1, 6) + " " + t)
            elif "list" in style:
                out.append(f"- {t}")
            else:
                out.append(t)
        elif tag == "tbl":
            tbl = docx.table.Table(block, d)
            rows = [" | ".join(c.text.strip().replace("\n", " ") for c in r.cells) for r in tbl.rows]
            if rows:
                ncol = rows[0].count("|") + 1
                out.append("\n".join([f"| {rows[0]} |", "|" + "---|" * ncol] + [f"| {r} |" for r in rows[1:]]))
    return "\n\n".join(out)


def pdf_md(p: Path) -> str:
    from pypdf import PdfReader
    r = PdfReader(str(p))
    return "\n\n".join(f"## Page {i}\n\n{(pg.extract_text() or '').strip()}" for i, pg in enumerate(r.pages, 1))


def extract(p: Path) -> str:
    ext = p.suffix.lower()
    if ext == ".pptx":
        return pptx_md(p)
    if ext == ".docx":
        return docx_md(p)
    if ext == ".pdf":
        return pdf_md(p)
    return p.read_text(encoding="utf-8", errors="replace").strip()


for item in json.loads(Path(sys.argv[1]).read_text(encoding="utf-8")):
    src, out = Path(item["src"]), Path(item["out"])
    if out.exists():
        raise SystemExit(f"refusing to overwrite {out}")
    body = extract(src)
    words = re.sub(r"\s+", " ", re.sub(r"[#*|>`-]", " ", body)).strip()
    summary = item.get("summary") or (words[:240].rsplit(" ", 1)[0] + "…" if len(words) > 240 else words)
    source = item.get("source") or src.as_posix().replace(HOME, "~")
    fm = [
        "---",
        f"title: {q(item['title'])}",
        f"source: {q(source)}",
        f"type: {item.get('type', 'notes')}",
        f"ingested: {TODAY}",
        f"tags: [{', '.join(item.get('tags', []))}]",
        f"summary: {q(summary)}",
        f"authors: [{', '.join(item.get('authors', ['unknown']))}]",
        f"source_date: {item.get('source_date', 'unknown')}",
    ]
    for k, v in (item.get("extra") or {}).items():
        fm.append(f"{k}: {v}")
    fm += ["---", "", f"# {item['title']}", "",
           f"> Text extracted {TODAY} from `{src.name}` ({src.suffix.lstrip('.')}). Original: `{source}`.",
           "", body, ""]
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(fm), encoding="utf-8")
    print(f"{len(body):>7} chars  {out.name}")
